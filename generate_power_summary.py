"""
generate_power_summary.py — AMPIN_FTG_4 Power Summary Generator (SFTP)
=========================================================================
Connects directly to the destination SFTP server (same account as
ftp_to_sftp.py) and reads the raw FTP_REPORT .xls files straight out of
the Green 8 / Green 3 / Green 10 folders - no local copies needed.

Produces one Excel workbook with one sheet per plant, each containing
just:

    Date | Time Block | Total Solar Active Power (MW) | Total Wind Active
    Power (MW) | Total Hybrid Active Power (MW)

COLUMN MAPPING (confirmed against a real Green 8/SECI-3 file; Green 3/C&I
and Green 10/SECI-5 use the mapping already established for the dashboard)
--------------------------------------------------------------------------
  Green 8  (SECI-3): Timestamp=A  Hybrid=B  Solar=CQ+CU (FDR-12+FDR-13)   Wind=CO (Total WTG Active Power)
  Green 3  (C&I):    Timestamp=A  Hybrid=B  Solar=BH                     Wind=BF
  Green 10 (SECI-5): Timestamp=B  Hybrid=C  Solar=CM+CQ                  Wind=CK

Each FTP_REPORT file already contains every 15-minute row from the start
of its day up through the file's own timestamp, so for a given date we
only need the LATEST file for that date on the server - it already has
the whole day.

Each row's timestamp is shifted back by 15 minutes for the "Time Block"
column, matching the existing dashboard's convention.

SETUP
-----
Put this file next to your existing ".env" (the same one ftp_to_sftp.py
uses - DEST_SFTP_HOST, DEST_SFTP_USER, DEST_SFTP_KEY_PATH,
DEST_BASE_PATH, etc.). No new settings are needed.

    pip install pandas openpyxl xlrd paramiko python-dotenv --break-system-packages
    python generate_power_summary.py

Output is written to power_summary.xlsx in this same folder.
"""

import os
import re
import io
import sys
import stat
from pathlib import Path, PurePosixPath
from datetime import datetime, timedelta

import pandas as pd
import openpyxl.utils as xl_utils
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils.dataframe import dataframe_to_rows

try:
    from dotenv import load_dotenv
except ImportError:
    print("Missing dependency. Run:  pip install python-dotenv")
    sys.exit(1)

try:
    import paramiko
except ImportError:
    print("Missing dependency. Run:  pip install paramiko")
    sys.exit(1)

SCRIPT_DIR = Path(__file__).resolve().parent
ENV_PATH = SCRIPT_DIR / ".env"
if ENV_PATH.exists():
    load_dotenv(ENV_PATH)


def env(key, default=None, required=True):
    val = os.getenv(key, default)
    if required and (val is None or val == ""):
        print(f"Missing required setting '{key}' in .env file.")
        sys.exit(1)
    return val


DEST_HOST = env("DEST_SFTP_HOST")
DEST_PORT = int(env("DEST_SFTP_PORT", "22"))
DEST_USER = env("DEST_SFTP_USER")
DEST_PASS = env("DEST_SFTP_PASS", "", required=False)
DEST_KEY_PATH = env("DEST_SFTP_KEY_PATH", "", required=False)
DEST_KEY_PASSPHRASE = env("DEST_SFTP_KEY_PASSPHRASE", "", required=False) or None
DEST_BASE_PATH = env("DEST_BASE_PATH", "/srv/data/sftp-users/indian-operations/ampin")

OUTPUT_PATH = SCRIPT_DIR / "power_summary.xlsx"

PLANTS = {
    "Green 8": {
        "timestamp": "A",
        "hybrid": "B",
        "solar": ["CQ", "CU"],
        "wind": "CO",
    },
    "Green 3": {
        "timestamp": "A",
        "hybrid": "B",
        "solar": "BH",
        "wind": "BF",
    },
    "Green 10": {
        "timestamp": "B",
        "hybrid": "C",
        "solar": ["CM", "CQ"],
        "wind": "CK",
    },
}

TIMESTAMP_DISPLAY_OFFSET_MINUTES = 15

FILENAME_PATTERN = re.compile(r"(\d{2}-\d{2}-\d{4})[ _](\d{1,2})_(\d{2})_(\d{2})")


# --------------------------------------------------------------------------
# SFTP connection (same approach as ftp_to_sftp.py)
# --------------------------------------------------------------------------

def load_private_key(path: str, passphrase):
    key_classes = [paramiko.RSAKey, paramiko.Ed25519Key, paramiko.ECDSAKey]
    last_error = None
    for key_class in key_classes:
        try:
            return key_class.from_private_key_file(path, password=passphrase)
        except Exception as e:
            last_error = e
    raise RuntimeError(
        f"Could not parse private key at {path} with any known key type "
        f"(tried RSA/Ed25519/ECDSA). Last error: {last_error}"
    )


def sftp_connect():
    ssh = paramiko.SSHClient()
    ssh.set_missing_host_key_policy(paramiko.AutoAddPolicy())

    if DEST_KEY_PATH:
        pkey = load_private_key(DEST_KEY_PATH, DEST_KEY_PASSPHRASE)
        ssh.connect(
            DEST_HOST,
            port=DEST_PORT,
            username=DEST_USER,
            pkey=pkey,
            allow_agent=False,
            look_for_keys=False,
            timeout=30,
        )
    else:
        ssh.connect(
            DEST_HOST,
            port=DEST_PORT,
            username=DEST_USER,
            password=DEST_PASS,
            allow_agent=False,
            look_for_keys=False,
            timeout=30,
        )

    sftp = ssh.open_sftp()
    return sftp, ssh


# --------------------------------------------------------------------------
# Filename parsing (same convention as the dashboard/relay scripts)
# --------------------------------------------------------------------------

def parse_filename(filename: str):
    m = FILENAME_PATTERN.search(filename)
    if not m:
        return None
    date_str, hh, mm, ss = m.groups()
    try:
        d = datetime.strptime(date_str, "%d-%m-%Y").date()
    except ValueError:
        return None
    return d, (int(hh), int(mm), int(ss))


def latest_remote_file_per_date(sftp: paramiko.SFTPClient, remote_dir: str):
    """Returns {date: filename} keeping only the latest-timestamped file per date."""
    try:
        entries = sftp.listdir_attr(remote_dir)
    except FileNotFoundError:
        return {}

    best = {}
    for entry in entries:
        if stat.S_ISDIR(entry.st_mode):
            continue
        if not entry.filename.lower().endswith((".xls", ".xlsx")):
            continue
        parsed = parse_filename(entry.filename)
        if not parsed:
            continue
        d, t = parsed
        if d not in best or t > best[d][0]:
            best[d] = (t, entry.filename)
    return {d: fname for d, (t, fname) in best.items()}


def download_to_memory(sftp: paramiko.SFTPClient, remote_path: str) -> io.BytesIO:
    buf = io.BytesIO()
    sftp.getfo(remote_path, buf)
    buf.seek(0)
    return buf


# --------------------------------------------------------------------------
# Excel parsing (column-letter based, mirrors dashboard.py's approach)
# --------------------------------------------------------------------------

def _engine_for(filename: str):
    return "xlrd" if filename.lower().endswith(".xls") else "openpyxl"


def _find_header_row(file_bytes: io.BytesIO, engine: str) -> int:
    file_bytes.seek(0)
    preview = pd.read_excel(file_bytes, sheet_name=0, header=None, nrows=6, engine=engine)
    for idx, row in preview.iterrows():
        row_vals = [str(v).strip().upper() for v in row.tolist()]
        if any("TIMESTAMP" in v for v in row_vals):
            return idx
    return 1


def _col_index(letter: str) -> int:
    return xl_utils.column_index_from_string(letter) - 1


def _sum_columns(df: pd.DataFrame, letters):
    if isinstance(letters, str):
        letters = [letters]
    total = None
    for letter in letters:
        idx = _col_index(letter)
        if idx >= df.shape[1]:
            raise ValueError(
                f"Column '{letter}' (position {idx}) is out of range - "
                f"this file only has {df.shape[1]} columns."
            )
        series = pd.to_numeric(df.iloc[:, idx], errors="coerce")
        total = series if total is None else (total + series)
    return total


def parse_report_bytes(file_bytes: io.BytesIO, filename: str, columns: dict) -> pd.DataFrame:
    """Returns a DataFrame with: date, time_block, solar_mw, wind_mw, hybrid_mw."""
    engine = _engine_for(filename)
    header_row = _find_header_row(file_bytes, engine)
    file_bytes.seek(0)
    df = pd.read_excel(file_bytes, sheet_name=0, header=None, engine=engine)
    df = df.iloc[header_row + 1:].reset_index(drop=True)

    ts_idx = _col_index(columns["timestamp"])
    timestamps = pd.to_datetime(df.iloc[:, ts_idx], errors="coerce")

    hybrid_idx = _col_index(columns["hybrid"])
    hybrid = pd.to_numeric(df.iloc[:, hybrid_idx], errors="coerce")

    solar = _sum_columns(df, columns["solar"])
    wind = _sum_columns(df, columns["wind"])

    clean = pd.DataFrame({
        "timestamp": timestamps,
        "solar_mw": solar,
        "wind_mw": wind,
        "hybrid_mw": hybrid,
    })
    clean = clean.dropna(subset=["timestamp"]).sort_values("timestamp")
    clean["display_timestamp"] = clean["timestamp"] - timedelta(minutes=TIMESTAMP_DISPLAY_OFFSET_MINUTES)
    clean["date"] = clean["display_timestamp"].dt.strftime("%d-%m-%Y")
    clean["time_block"] = clean["display_timestamp"].dt.strftime("%H:%M")
    return clean[["date", "time_block", "solar_mw", "wind_mw", "hybrid_mw", "display_timestamp"]]


def build_plant_table(plant_name: str, columns: dict, sftp: paramiko.SFTPClient, remote_dir: str) -> pd.DataFrame:
    files_by_date = latest_remote_file_per_date(sftp, remote_dir)
    if not files_by_date:
        print(f"[{plant_name}] no matching FTP_REPORT files found in {remote_dir}")
        return pd.DataFrame(columns=["date", "time_block", "solar_mw", "wind_mw", "hybrid_mw"])

    frames = []
    for d, filename in sorted(files_by_date.items()):
        remote_path = str(PurePosixPath(remote_dir) / filename)
        try:
            file_bytes = download_to_memory(sftp, remote_path)
            frames.append(parse_report_bytes(file_bytes, filename, columns))
        except Exception as e:
            print(f"[{plant_name}] failed to parse {filename}: {e}")

    if not frames:
        return pd.DataFrame(columns=["date", "time_block", "solar_mw", "wind_mw", "hybrid_mw"])

    combined = pd.concat(frames, ignore_index=True)
    combined = combined.drop_duplicates(subset=["display_timestamp"]).sort_values("display_timestamp")
    print(f"[{plant_name}] {len(files_by_date)} day(s), {len(combined)} time-block rows")
    return combined[["date", "time_block", "solar_mw", "wind_mw", "hybrid_mw"]]


# --------------------------------------------------------------------------
# Write output workbook
# --------------------------------------------------------------------------

def write_workbook(tables: dict, output_path: Path):
    wb = Workbook()
    wb.remove(wb.active)

    header_font = Font(name="Arial", bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", start_color="2F3A56")
    header_align = Alignment(horizontal="center")
    body_font = Font(name="Arial")

    headers = [
        "Date",
        "Time Block",
        "Total Solar Active Power (MW)",
        "Total Wind Active Power (MW)",
        "Total Hybrid Active Power (MW)",
    ]

    for plant_name, df in tables.items():
        ws = wb.create_sheet(title=plant_name[:31])

        for col_idx, h in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align

        out_df = df.rename(columns={
            "date": "Date",
            "time_block": "Time Block",
            "solar_mw": "Total Solar Active Power (MW)",
            "wind_mw": "Total Wind Active Power (MW)",
            "hybrid_mw": "Total Hybrid Active Power (MW)",
        })
        out_df = out_df.round(3)

        for row in dataframe_to_rows(out_df, index=False, header=False):
            ws.append(row)
            for cell in ws[ws.max_row]:
                cell.font = body_font

        ws.column_dimensions["A"].width = 14
        ws.column_dimensions["B"].width = 12
        for letter in ["C", "D", "E"]:
            ws.column_dimensions[letter].width = 28

        if len(out_df) == 0:
            ws.cell(row=2, column=1, value="No data found for this plant.")

    if not wb.sheetnames:
        wb.create_sheet(title="No Data")

    wb.save(output_path)


def main():
    print(f"Connecting to {DEST_HOST}:{DEST_PORT} as {DEST_USER}...")
    try:
        sftp, ssh = sftp_connect()
    except Exception as e:
        print(f"Could not connect to destination SFTP server: {type(e).__name__}: {e}")
        sys.exit(1)

    try:
        tables = {}
        for plant_name, columns in PLANTS.items():
            remote_dir = str(PurePosixPath(DEST_BASE_PATH) / plant_name)
            tables[plant_name] = build_plant_table(plant_name, columns, sftp, remote_dir)
    finally:
        try:
            sftp.close()
            ssh.close()
        except Exception:
            pass

    write_workbook(tables, OUTPUT_PATH)
    print(f"\nWrote: {OUTPUT_PATH}")


if __name__ == "__main__":
    main()