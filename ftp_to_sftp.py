r"""
ftp_to_sftp.py — AMPIN_FTG_4 FTP -> SFTP Relay (trimmed columns)
=================================================================
Polls the three source FTP servers (SECI-3, C&I, SECI-5). For each raw
report file, this script PARSES it and uploads a small trimmed .xlsx
containing only:

    Date | Time Block | Total Solar Active Power (MW) |
    Total Wind Active Power (MW) | Total Hybrid Active Power (MW)

...into the matching "Green N" folder under a shared base path:

    SECI-3  ->  <DEST_BASE_PATH>/Green 8
    C&I     ->  <DEST_BASE_PATH>/Green 3
    SECI-5  ->  <DEST_BASE_PATH>/Green 10

COLUMN MAPPING
--------------
  SECI-3 (Green 8):  Timestamp=A  Hybrid=B  Solar=CQ+CU  Wind=CO
  C&I    (Green 3):  Timestamp=A  Hybrid=B  Solar=BH     Wind=BF
  SECI-5 (Green 10): Timestamp=B  Hybrid=C  Solar=CM+CQ  Wind=CK

Each row's timestamp is shifted back by 15 minutes for the "Time Block"
column, matching the existing dashboard's convention.

MSEDCL is intentionally NOT part of this relay.

Only files whose date (parsed from the filename) is on/after
SYNC_START_DATE are considered (default 2026-09-01).

Before each poll run, the script lists what's already in the destination
folder and only processes raw files that don't yet have a trimmed
counterpart (same basename, .xlsx extension).

SETUP
-----
1. Keep this file next to your ".env" (with SECI3_HOST etc.) and add:

       DEST_SFTP_HOST=sftp-emsys-india.energymeteo.systems
       DEST_SFTP_PORT=22
       DEST_SFTP_USER=indian-operations
       DEST_BASE_PATH=/srv/data/sftp-users/indian-operations/ampin
       SYNC_START_DATE=2026-09-01
       POLL_INTERVAL_MINUTES=5

   Then ONE of:
     a) Key auth:
            DEST_SFTP_KEY_PATH=C:\path\to\key.pem
            DEST_SFTP_KEY_PASSPHRASE=
     b) Password auth:
            DEST_SFTP_PASS=your-password-here

2. Install dependencies (pip.exe is blocked, so go through python):
       python -m pip install --user python-dotenv paramiko pandas openpyxl xlrd

3. Run:
       python ftp_to_sftp.py                 (asks which plants to sync)
       python ftp_to_sftp.py green8 green10  (no prompt)
       python ftp_to_sftp.py all
"""

import os
import re
import io
import sys
import time
import stat
from pathlib import Path, PurePosixPath
from datetime import datetime, date, timedelta
from ftplib import FTP, error_perm

try:
    from dotenv import load_dotenv
except ImportError:
    print("Missing dependency. Run:  python -m pip install --user python-dotenv")
    sys.exit(1)

try:
    import paramiko
except ImportError:
    print("Missing dependency. Run:  python -m pip install --user paramiko")
    sys.exit(1)

try:
    import pandas as pd
    import openpyxl.utils as xl_utils
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils.dataframe import dataframe_to_rows
except ImportError:
    print("Missing dependency. Run:  python -m pip install --user pandas openpyxl xlrd")
    sys.exit(1)

SCRIPT_DIR = Path(__file__).resolve().parent
ENV_PATH = SCRIPT_DIR / ".env"
if ENV_PATH.exists():
    load_dotenv(ENV_PATH)


def env(key, default=None, required=True):
    val = os.getenv(key, default)
    if required and (val is None or val == ""):
        print(f"Missing required setting '{key}' in .env file.")
        sys.exit(1)
    return val


if not ENV_PATH.exists() and not os.getenv("SECI3_HOST"):
    print(
        "No FTP credentials found. Copy your .env (with SECI3_HOST etc.) "
        "into the same folder as this script."
    )
    sys.exit(1)


SOURCES = {
    "SECI-3": {
        "host": env("SECI3_HOST"),
        "port": int(env("SECI3_PORT", "21")),
        "user": env("SECI3_USER"),
        "password": env("SECI3_PASS"),
        "remote_dir": env("SECI3_REMOTE_DIR", "/", required=False) or "/",
    },
    "CNI": {
        "host": env("CNI_HOST"),
        "port": int(env("CNI_PORT", "21")),
        "user": env("CNI_USER"),
        "password": env("CNI_PASS"),
        "remote_dir": env("CNI_REMOTE_DIR", "/", required=False) or "/",
    },
    "SECI-5": {
        "host": env("SECI5_HOST"),
        "port": int(env("SECI5_PORT", "21")),
        "user": env("SECI5_USER"),
        "password": env("SECI5_PASS"),
        "remote_dir": env("SECI5_REMOTE_DIR", "/", required=False) or "/",
    },
}

# Which destination "Green N" folder each source relays into.
SOURCE_TO_DEST_FOLDER = {
    "SECI-3": "Green 8",
    "CNI": "Green 3",
    "SECI-5": "Green 10",
}

# Column mapping per source, used to build the trimmed output.
SOURCE_COLUMNS = {
    "SECI-3": {"timestamp": "A", "hybrid": "B", "solar": ["CQ", "CU"], "wind": "CO"},
    "CNI": {"timestamp": "A", "hybrid": "B", "solar": "BH", "wind": "BF"},
    "SECI-5": {"timestamp": "B", "hybrid": "C", "solar": ["CM", "CQ"], "wind": "CK"},
}

TIMESTAMP_DISPLAY_OFFSET_MINUTES = 15

DEST_HOST = env("DEST_SFTP_HOST")
DEST_PORT = int(env("DEST_SFTP_PORT", "22"))
DEST_USER = env("DEST_SFTP_USER")
DEST_PASS = env("DEST_SFTP_PASS", "", required=False)
DEST_KEY_PATH = env("DEST_SFTP_KEY_PATH", "", required=False)
DEST_KEY_PASSPHRASE = env("DEST_SFTP_KEY_PASSPHRASE", "", required=False) or None
DEST_BASE_PATH = env("DEST_BASE_PATH", "/srv/data/sftp-users/indian-operations/ampin")

POLL_INTERVAL_MINUTES = float(env("POLL_INTERVAL_MINUTES", "5", required=False) or 5)

SYNC_START_DATE_STR = env("SYNC_START_DATE", "2026-09-01", required=False)
try:
    SYNC_START_DATE = datetime.strptime(SYNC_START_DATE_STR, "%Y-%m-%d").date()
except ValueError:
    print(f"SYNC_START_DATE '{SYNC_START_DATE_STR}' is not in YYYY-MM-DD format, defaulting to 2026-09-01")
    SYNC_START_DATE = date(2026, 9, 1)

FILENAME_PATTERN = re.compile(r"(\d{2}-\d{2}-\d{4})[ _](\d{1,2})_(\d{2})_(\d{2})")


def parse_xls_date(fname: str):
    m = FILENAME_PATTERN.search(fname)
    if not m:
        return None
    try:
        return datetime.strptime(m.group(1), "%d-%m-%Y").date()
    except ValueError:
        return None


def trimmed_filename(raw_basename: str) -> str:
    """Same basename as the raw file, but with a .xlsx extension."""
    return str(Path(raw_basename).with_suffix(".xlsx"))


def ftp_connect(source: dict) -> FTP:
    ftp = FTP()
    ftp.connect(source["host"], source["port"], timeout=30)
    ftp.login(source["user"], source["password"])
    ftp.set_pasv(True)
    ftp.cwd(source["remote_dir"])
    return ftp


def load_private_key(path: str, passphrase):
    """Try each paramiko key class in turn and return the first that parses."""
    key_classes = [paramiko.RSAKey, paramiko.Ed25519Key, paramiko.ECDSAKey]
    last_error = None
    for key_class in key_classes:
        try:
            return key_class.from_private_key_file(path, password=passphrase)
        except Exception as e:
            last_error = e
    raise RuntimeError(
        f"Could not parse private key at {path} with any known key type "
        f"(tried RSA/Ed25519/ECDSA). Last error: {last_error}"
    )


def sftp_connect():
    """Connect with an SSH key if DEST_SFTP_KEY_PATH is set, else password."""
    ssh = paramiko.SSHClient()
    ssh.set_missing_host_key_policy(paramiko.AutoAddPolicy())

    if DEST_KEY_PATH:
        pkey = load_private_key(DEST_KEY_PATH, DEST_KEY_PASSPHRASE)
        try:
            ssh.connect(
                DEST_HOST,
                port=DEST_PORT,
                username=DEST_USER,
                pkey=pkey,
                allow_agent=False,
                look_for_keys=False,
                timeout=30,
            )
        except paramiko.ssh_exception.AuthenticationException:
            # Some servers don't advertise supported RSA signature algorithms
            # correctly. Retry once, forcing the classic ssh-rsa (SHA-1).
            ssh.connect(
                DEST_HOST,
                port=DEST_PORT,
                username=DEST_USER,
                pkey=pkey,
                allow_agent=False,
                look_for_keys=False,
                timeout=30,
                disabled_algorithms={"pubkeys": ["rsa-sha2-256", "rsa-sha2-512"]},
            )
    else:
        ssh.connect(
            DEST_HOST,
            port=DEST_PORT,
            username=DEST_USER,
            password=DEST_PASS,
            allow_agent=False,
            look_for_keys=False,
            timeout=30,
        )

    sftp = ssh.open_sftp()
    return sftp, ssh


def ensure_remote_dir(sftp: paramiko.SFTPClient, remote_dir: str):
    parts = PurePosixPath(remote_dir).parts
    current = ""
    for part in parts:
        current = part if current == "" else str(PurePosixPath(current) / part)
        try:
            sftp.stat(current)
        except FileNotFoundError:
            try:
                sftp.mkdir(current)
            except Exception as e:
                print(f"  warning: could not create remote dir {current}: {e}")


def list_remote_filenames(sftp: paramiko.SFTPClient, remote_dir: str):
    try:
        entries = sftp.listdir_attr(remote_dir)
    except FileNotFoundError:
        return set()
    return {e.filename for e in entries if not stat.S_ISDIR(e.st_mode)}


# --------------------------------------------------------------------------
# Transform: raw xls bytes -> trimmed xlsx bytes
# --------------------------------------------------------------------------

def _engine_for(filename: str):
    return "xlrd" if filename.lower().endswith(".xls") else "openpyxl"


def _find_header_row(file_bytes: io.BytesIO, engine: str) -> int:
    file_bytes.seek(0)
    preview = pd.read_excel(file_bytes, sheet_name=0, header=None, nrows=6, engine=engine)
    for idx, row in preview.iterrows():
        row_vals = [str(v).strip().upper() for v in row.tolist()]
        if any("TIMESTAMP" in v for v in row_vals):
            return idx
    return 1


def _col_index(letter: str) -> int:
    return xl_utils.column_index_from_string(letter) - 1


def _sum_columns(df: pd.DataFrame, letters):
    if isinstance(letters, str):
        letters = [letters]
    total = None
    for letter in letters:
        idx = _col_index(letter)
        if idx >= df.shape[1]:
            raise ValueError(
                f"Column '{letter}' (position {idx}) is out of range - "
                f"this file only has {df.shape[1]} columns."
            )
        series = pd.to_numeric(df.iloc[:, idx], errors="coerce")
        total = series if total is None else (total + series)
    return total


def parse_report_bytes(file_bytes: io.BytesIO, filename: str, columns: dict) -> pd.DataFrame:
    """Returns a DataFrame with: Date, Time Block, solar, wind, hybrid columns."""
    engine = _engine_for(filename)
    header_row = _find_header_row(file_bytes, engine)
    file_bytes.seek(0)
    df = pd.read_excel(file_bytes, sheet_name=0, header=None, engine=engine)
    df = df.iloc[header_row + 1:].reset_index(drop=True)

    ts_idx = _col_index(columns["timestamp"])
    timestamps = pd.to_datetime(df.iloc[:, ts_idx], errors="coerce")

    hybrid_idx = _col_index(columns["hybrid"])
    hybrid = pd.to_numeric(df.iloc[:, hybrid_idx], errors="coerce")

    solar = _sum_columns(df, columns["solar"])
    wind = _sum_columns(df, columns["wind"])

    clean = pd.DataFrame({
        "timestamp": timestamps,
        "solar_mw": solar,
        "wind_mw": wind,
        "hybrid_mw": hybrid,
    })
    clean = clean.dropna(subset=["timestamp"]).sort_values("timestamp")
    clean["display_timestamp"] = clean["timestamp"] - timedelta(minutes=TIMESTAMP_DISPLAY_OFFSET_MINUTES)
    clean["Date"] = clean["display_timestamp"].dt.strftime("%d-%m-%Y")
    clean["Time Block"] = clean["display_timestamp"].dt.strftime("%H:%M")
    out = clean.rename(columns={
        "solar_mw": "Total Solar Active Power (MW)",
        "wind_mw": "Total Wind Active Power (MW)",
        "hybrid_mw": "Total Hybrid Active Power (MW)",
    })
    return out[[
        "Date", "Time Block",
        "Total Solar Active Power (MW)",
        "Total Wind Active Power (MW)",
        "Total Hybrid Active Power (MW)",
    ]].round(3)


def build_trimmed_xlsx_bytes(df: pd.DataFrame) -> io.BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "Power Summary"

    header_font = Font(name="Arial", bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", start_color="2F3A56")
    header_align = Alignment(horizontal="center")
    body_font = Font(name="Arial")

    for col_idx, h in enumerate(df.columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align

    for row in dataframe_to_rows(df, index=False, header=False):
        ws.append(row)
        for cell in ws[ws.max_row]:
            cell.font = body_font

    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 12
    for letter in ["C", "D", "E"]:
        ws.column_dimensions[letter].width = 28

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# --------------------------------------------------------------------------
# Relay
# --------------------------------------------------------------------------

def relay_source(source_name: str, sftp: paramiko.SFTPClient):
    source = SOURCES[source_name]
    columns = SOURCE_COLUMNS[source_name]
    dest_folder = SOURCE_TO_DEST_FOLDER[source_name]
    remote_dest_dir = str(PurePosixPath(DEST_BASE_PATH) / dest_folder)
    ensure_remote_dir(sftp, remote_dest_dir)

    already_uploaded = list_remote_filenames(sftp, remote_dest_dir)

    try:
        ftp = ftp_connect(source)
    except Exception as e:
        print(f"[{source_name}] connect/login to source FTP failed: {e}")
        return

    relayed, errors = 0, 0
    try:
        try:
            names = ftp.nlst()
        except error_perm as e:
            print(f"[{source_name}] cannot list source directory: {e}")
            return

        matching_files = []
        for remote_name in names:
            if not remote_name.lower().endswith((".xls", ".xlsx")):
                continue
            basename = remote_name.split("/")[-1]
            fdate = parse_xls_date(basename)
            if fdate is not None and fdate >= SYNC_START_DATE:
                matching_files.append(remote_name)

        to_process = [
            n for n in matching_files
            if trimmed_filename(n.split("/")[-1]) not in already_uploaded
        ]
        skipped = len(matching_files) - len(to_process)
        print(
            f"[{source_name}] {len(matching_files)} file(s) match date range, "
            f"{len(to_process)} still need processing, {skipped} already there",
            flush=True,
        )

        for remote_name in to_process:
            basename = remote_name.split("/")[-1]
            out_name = trimmed_filename(basename)
            try:
                raw_buf = io.BytesIO()
                ftp.retrbinary(f"RETR {remote_name}", raw_buf.write)
                raw_buf.seek(0)

                df = parse_report_bytes(raw_buf, basename, columns)
                trimmed_buf = build_trimmed_xlsx_bytes(df)

                dest_path = str(PurePosixPath(remote_dest_dir) / out_name)
                sftp.putfo(trimmed_buf, dest_path)
                relayed += 1
                print(f"[{source_name} -> {dest_folder}] uploaded {out_name} ({relayed} so far)", flush=True)
            except Exception as e:
                errors += 1
                print(f"[{source_name}] failed to process {basename}: {e}", flush=True)
    finally:
        try:
            ftp.quit()
        except Exception:
            ftp.close()

    print(f"[{source_name} -> {dest_folder}] uploaded={relayed} skipped(existing)={skipped} errors={errors}")


def sync_all(active_sources):
    print(f"\n=== Relay run started {datetime.now().strftime('%Y-%m-%d %H:%M:%S')} (from {SYNC_START_DATE}) ===")
    try:
        sftp, ssh = sftp_connect()
    except Exception as e:
        print(f"Could not connect to destination SFTP server: {type(e).__name__}: {e}")
        print(
            "  -> Try logging into the same host/port/user (+ key or password) "
            "with FileZilla or WinSCP to confirm the credentials work outside "
            "this script."
        )
        return

    try:
        for src in active_sources:
            relay_source(src, sftp)
    finally:
        try:
            sftp.close()
            ssh.close()
        except Exception:
            pass


# --------------------------------------------------------------------------
# Source selection (interactive menu or command-line args)
# --------------------------------------------------------------------------

CHOICE_MAP = {
    "1": "SECI-3", "green 8": "SECI-3", "green8": "SECI-3", "seci-3": "SECI-3", "seci3": "SECI-3",
    "2": "CNI", "green 3": "CNI", "green3": "CNI", "cni": "CNI", "c&i": "CNI",
    "3": "SECI-5", "green 10": "SECI-5", "green10": "SECI-5", "seci-5": "SECI-5", "seci5": "SECI-5",
}


def resolve_choice(raw: str):
    key = raw.strip().lower()
    if key in ("4", "all", "a"):
        return list(SOURCES.keys())
    return [CHOICE_MAP[key]] if key in CHOICE_MAP else None


def prompt_for_sources():
    print("\nWhich plant(s) do you want to sync?")
    print("  1) Green 8  (SECI-3)")
    print("  2) Green 3  (C&I)")
    print("  3) Green 10 (SECI-5)")
    print("  4) All three")
    while True:
        raw = input("Enter one or more, comma-separated (e.g. 1,3), or 'all': ").strip()
        if not raw:
            continue
        chosen = set()
        ok = True
        for part in raw.split(","):
            resolved = resolve_choice(part)
            if resolved is None:
                print(f"  Didn't recognize '{part.strip()}' - try again.")
                ok = False
                break
            chosen.update(resolved)
        if ok and chosen:
            return [s for s in SOURCES if s in chosen]


def get_active_sources():
    """
    CLI args (e.g. `python ftp_to_sftp.py green8 green10`) select sources
    without prompting, so this can be scheduled/automated. With no args,
    prompts interactively once at startup.
    """
    if len(sys.argv) > 1:
        chosen = set()
        for arg in sys.argv[1:]:
            resolved = resolve_choice(arg)
            if resolved is None:
                print(f"Unrecognized source argument: '{arg}'")
                print("Valid options: 1/2/3/4, green8/green3/green10, seci-3/cni/seci-5, all")
                sys.exit(1)
            chosen.update(resolved)
        return [s for s in SOURCES if s in chosen]
    return prompt_for_sources()


def main():
    active_sources = get_active_sources()
    dest_folders = ", ".join(SOURCE_TO_DEST_FOLDER[s] for s in active_sources)
    print(f"\nSyncing: {dest_folders}")
    print(f"Relaying (trimmed) into: {DEST_HOST}:{DEST_PORT} -> {DEST_BASE_PATH}")
    print(f"Polling every {POLL_INTERVAL_MINUTES} minute(s). Press Ctrl+C to stop.")
    while True:
        try:
            sync_all(active_sources)
        except Exception as e:
            print(f"Unexpected error during sync: {e}")
        time.sleep(POLL_INTERVAL_MINUTES * 60)


if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        print("\nStopped.")